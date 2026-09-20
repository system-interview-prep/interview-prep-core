import io

import pytest
from openpyxl import load_workbook

from src.modules.question_bank.import_parser import (
    IMPORT_COLUMNS,
    csv_template,
    parse_csv,
    parse_xlsx,
    xlsx_template,
)


def valid_row() -> str:
    values = {key: "x" for key in IMPORT_COLUMNS}
    values.update({
        "thinking_seconds": "0", "soft_answer_seconds": "60", "hard_answer_seconds": "120",
        "primary_competency_id": "comp", "target_role_ids": "[]",
        "supporting_competency_ids": "[]", "skill_ids": "[]", "expected_points": "[]",
        "source_refs": "[]",
    })
    return ",".join(f'"{values[key]}"' for key in IMPORT_COLUMNS)


def test_template_has_canonical_header() -> None:
    assert csv_template().decode().splitlines()[0].split(",") == list(IMPORT_COLUMNS)


def test_csv_parser_reports_json_and_duration_errors() -> None:
    invalid = valid_row().replace('"120"', '"10"').replace('"[]"', '"bad-json"', 1)
    content = (",".join(IMPORT_COLUMNS) + "\n" + invalid).encode()
    errors = parse_csv(content)[0].errors
    assert {error["code"] for error in errors} == {"INVALID_JSON", "HARD_BELOW_SOFT"}


def test_csv_parser_rejects_wrong_header() -> None:
    with pytest.raises(ValueError, match="header"):
        parse_csv(b"stable_key\nfoo\n")


def test_xlsx_template_and_parser_round_trip() -> None:
    workbook = load_workbook(io.BytesIO(xlsx_template()))
    assert workbook.sheetnames == ["Questions", "Instructions"]
    worksheet = workbook["Questions"]
    worksheet.append([
        "key", "1", "vi-VN", "Question", "Objective", "CONCEPTUAL", "MEDIUM", 0, 60,
        120, "v1", "[]", "comp", "[]", "[]", "[]", "[]", "rubric", "",
    ])
    stream = io.BytesIO()
    workbook.save(stream)
    parsed = parse_xlsx(stream.getvalue())
    assert parsed[0].payload["stable_key"] == "key"
    assert parsed[0].errors == []
